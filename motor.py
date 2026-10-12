"""
motor.py - Motor de Auditoria Contínua (sem dependência do Streamlit).

Pode ser usado de 3 formas:
  1) pela interface:        streamlit run app.py
  2) por linha de comando:  python motor.py lancamentos.csv -o resultado.xlsx
  3) como ferramenta de um agente de IA / outro sistema: from motor import analisar_arquivo
"""
from __future__ import annotations

import argparse
import difflib
import hashlib
import io
import json
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta, timezone

import numpy as np
import pandas as pd
from sklearn.ensemble import IsolationForest

VERSAO = "2.0.0"

OBRIGATORIAS = ["ID_Transacao", "Data", "Usuario", "Conta_Contabil", "Valor"]
OPCIONAIS = ["Fornecedor", "Numero_NF", "Hora", "Tipo_Lancamento", "Aprovador"]

DIAS_PT = {0: "Segunda", 1: "Terça", 2: "Quarta", 3: "Quinta", 4: "Sexta", 5: "Sábado", 6: "Domingo"}

ROTULOS = {
    "Alerta_SoD": "Aprovador = lançador (SoD)",
    "Alerta_Duplicidade": "Duplicidade / fraude de voucher",
    "Alerta_Fracionamento": "Fracionamento para fugir da alçada",
    "Alerta_Alcada": "Valor logo abaixo da alçada",
    "Alerta_Reincidencia": "Reincidência sub-alçada (mesmo usuário)",
    "Alerta_Nao_Util": "Fim de semana / feriado",
    "Alerta_Outlier_Valor": "Valor atípico para a conta",
    "Alerta_Fora_Horario": "Manual fora do expediente",
    "Alerta_Valor_Redondo": "Valor redondo",
}
PESOS_PADRAO = {
    "Alerta_SoD": 50, "Alerta_Duplicidade": 50, "Alerta_Fracionamento": 45, "Alerta_Alcada": 35,
    "Alerta_Reincidencia": 20, "Alerta_Nao_Util": 20, "Alerta_Outlier_Valor": 15, "Alerta_Fora_Horario": 15, "Alerta_Valor_Redondo": 10,
}
NIVEIS = [("Crítico", 70), ("Alto", 50), ("Médio", 30)]  # limites mínimos; abaixo disso = Baixo
ACOES = {
    "Crítico": "Investigar imediatamente e reter evidências",
    "Alto": "Revisar documentação de suporte em até 5 dias úteis",
    "Médio": "Incluir na amostra de revisão periódica",
    "Baixo": "Sem ação",
}


@dataclass
class Parametros:
    alcada: float = 10000.0           # alçada de aprovação (R$)
    margem_pct: float = 5.0           # "logo abaixo": % abaixo da alçada
    janela_dup_dias: int = 5          # janela para duplicidade
    tol_valor: float = 0.05           # tolerância de valor na duplicidade (R$)
    janela_frac_dias: int = 5         # janela para fracionamento
    parcela_min_pct: float = 35.0     # parcela mínima (% da alçada) para compor fracionamento
    z_corte: float = 3.0              # outlier: z-score do log do valor dentro da conta
    hora_ini: int = 7                 # expediente
    hora_fim: int = 20
    redondo_min: float = 5000.0       # valor redondo (múltiplo de R$ 1.000) a partir de
    contaminacao: float = 0.03        # proporção de anomalias esperada pela IA
    peso_ia: float = 30.0             # pontos máximos da IA
    pesos: dict = field(default_factory=lambda: dict(PESOS_PADRAO))


def brl(v: float) -> str:
    return "R$ " + f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


# ---------------------------------------------------------------------------
# Calendário (feriados nacionais + pontos facultativos usuais)
# ---------------------------------------------------------------------------
def _pascoa(ano: int) -> date:
    a, b, c = ano % 19, ano // 100, ano % 100
    d, e = b // 4, b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = c // 4, c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    mes = (h + l - 7 * m + 114) // 31
    dia = ((h + l - 7 * m + 114) % 31) + 1
    return date(ano, mes, dia)


def feriados(anos) -> dict:
    fixos = {
        (1, 1): "Confraternização Universal", (4, 21): "Tiradentes", (5, 1): "Dia do Trabalho",
        (9, 7): "Independência", (10, 12): "N. Sra. Aparecida", (11, 2): "Finados",
        (11, 15): "Proclamação da República", (11, 20): "Consciência Negra", (12, 25): "Natal",
    }
    out = {}
    for ano in anos:
        for (m, d), nome in fixos.items():
            out[date(ano, m, d)] = nome
        p = _pascoa(ano)
        out[p - timedelta(days=48)] = "Carnaval (segunda)"
        out[p - timedelta(days=47)] = "Carnaval (terça)"
        out[p - timedelta(days=2)] = "Sexta-feira Santa"
        out[p + timedelta(days=60)] = "Corpus Christi"
    return out


# ---------------------------------------------------------------------------
# Leitura e validação
# ---------------------------------------------------------------------------
def ler_csv_bytes(raw: bytes) -> pd.DataFrame:
    """Lê CSV com separador (; , tab |) e codificação detectados automaticamente."""
    try:
        texto = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = raw.decode("latin-1")
    primeira = texto.splitlines()[0] if texto.strip() else ""
    sep = max([";", ",", "\t", "|"], key=primeira.count)
    return pd.read_csv(io.StringIO(texto), sep=sep, dtype=str, keep_default_na=False)


def _numero(s: pd.Series) -> pd.Series:
    """Converte '1.234,56', 'R$ 1234,56' ou '1234.56' em número."""
    s = s.astype(str).str.replace(r"[R$\s]", "", regex=True)
    br = s.str.contains(",", regex=False)
    s = s.where(~br, s.str.replace(".", "", regex=False).str.replace(",", ".", regex=False))
    return pd.to_numeric(s, errors="coerce")


def validar(df: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    avisos: list[str] = []
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]
    faltando = [c for c in OBRIGATORIAS if c not in df.columns]
    if faltando:
        raise ValueError(
            f"Colunas obrigatórias ausentes: {', '.join(faltando)}. "
            f"Obrigatórias: {', '.join(OBRIGATORIAS)}. Opcionais: {', '.join(OPCIONAIS)}."
        )

    df["Data"] = pd.to_datetime(df["Data"], dayfirst=True, errors="coerce", format="mixed")
    if not pd.api.types.is_numeric_dtype(df["Valor"]):
        df["Valor"] = _numero(df["Valor"])

    invalidas = df["Data"].isna() | df["Valor"].isna()
    if invalidas.any():
        avisos.append(f"{int(invalidas.sum())} linha(s) com data ou valor inválido foram descartadas.")
        df = df[~invalidas]
    if df.empty:
        raise ValueError("Nenhuma linha válida para analisar.")

    for c in ["ID_Transacao", "Usuario", "Conta_Contabil"] + [o for o in OPCIONAIS if o in df.columns]:
        df[c] = df[c].fillna("").astype(str).str.strip()

    if df["ID_Transacao"].duplicated().any():
        avisos.append("Há ID_Transacao repetido no arquivo; verifique a extração.")
    if (df["Valor"] <= 0).any():
        avisos.append("Há valores zero/negativos (estornos?); eles entram nas regras, mas são tratados como 0 na IA.")

    if "Hora" in df.columns:
        h = df["Hora"].str.extract(r"(\d{1,2}):(\d{2})")
        df["Hora_Num"] = pd.to_numeric(h[0]) + pd.to_numeric(h[1]) / 60
    if len(df) < 50:
        avisos.append("Base com menos de 50 linhas: a IA e o teste de Benford perdem confiabilidade.")
    return df.reset_index(drop=True), avisos


# ---------------------------------------------------------------------------
# Regras de auditoria
# ---------------------------------------------------------------------------
def _duplicidades(df: pd.DataFrame, p: Parametros) -> pd.DataFrame:
    tem_forn = "Fornecedor" in df.columns and (df["Fornecedor"] != "").any()
    tem_nf = "Numero_NF" in df.columns
    chave = "Fornecedor" if tem_forn else "Usuario"
    flag = pd.Series(False, index=df.index)
    motivo = pd.Series("", index=df.index, dtype=object)

    for nome, g in df.sort_values("Data").groupby(chave):
        if nome == "":
            continue
        idx, datas = g.index.to_list(), g["Data"].to_list()
        vals, ids = g["Valor"].to_list(), g["ID_Transacao"].to_list()
        nfs = g["Numero_NF"].to_list() if tem_nf else None
        for i in range(len(idx)):
            for j in range(i + 1, len(idx)):
                if (datas[j] - datas[i]).days > p.janela_dup_dias:
                    break  # ordenado por data
                if abs(vals[i] - vals[j]) <= p.tol_valor:
                    extra = ""
                    if tem_nf and nfs[i] and nfs[j]:
                        sim = difflib.SequenceMatcher(None, nfs[i], nfs[j]).ratio()
                        if nfs[i] == nfs[j]:
                            extra = " - mesma NF"
                        elif sim >= 0.8:
                            extra = f" - NF com variação de digitação ({sim:.0%} similar)"
                    flag.loc[idx[i]] = flag.loc[idx[j]] = True
                    motivo.loc[idx[i]] += f"par com ID {ids[j]}{extra}; "
                    motivo.loc[idx[j]] += f"par com ID {ids[i]}{extra}; "

    # mesma NF do mesmo fornecedor lançada mais de uma vez (mesmo fora da janela)
    if tem_forn and tem_nf:
        rep = (df["Fornecedor"] != "") & (df["Numero_NF"] != "") & df.duplicated(["Fornecedor", "Numero_NF"], keep=False)
        for i in df.index[rep & (motivo == "")]:
            flag.loc[i] = True
            motivo.loc[i] = "mesma NF lançada mais de uma vez; "
    return pd.DataFrame({"Alerta_Duplicidade": flag, "Detalhe_Duplicidade": motivo.str.rstrip("; ")})


def _fracionamento(df: pd.DataFrame, p: Parametros) -> pd.DataFrame:
    """Vários lançamentos do mesmo usuário p/ o mesmo fornecedor, cada um abaixo da alçada,
    que somados em poucos dias a ultrapassam."""
    tem_forn = "Fornecedor" in df.columns and (df["Fornecedor"] != "").any()
    chave2 = "Fornecedor" if tem_forn else "Conta_Contabil"
    flag = pd.Series(False, index=df.index)
    motivo = pd.Series("", index=df.index, dtype=object)
    pmin = p.alcada * p.parcela_min_pct / 100
    cand = df[(df["Valor"] < p.alcada) & (df["Valor"] >= pmin)].sort_values("Data")

    for (_, k2), g in cand.groupby(["Usuario", chave2]):
        if len(g) < 2 or k2 == "":
            continue
        idx, datas, vals = g.index.to_list(), g["Data"].to_list(), g["Valor"].to_list()
        for i in range(len(idx)):
            j, soma = i, 0.0
            while j < len(idx) and (datas[j] - datas[i]).days <= p.janela_frac_dias:
                soma += vals[j]
                j += 1
            if j - i >= 2 and soma >= p.alcada:
                txt = f"{j - i} lançamentos somam {brl(soma)} em até {p.janela_frac_dias} dias, cada um abaixo da alçada"
                for k in range(i, j):
                    flag.loc[idx[k]] = True
                    motivo.loc[idx[k]] = txt
    return pd.DataFrame({"Alerta_Fracionamento": flag, "Detalhe_Fracionamento": motivo})


def aplicar_regras(df: pd.DataFrame, p: Parametros) -> pd.DataFrame:
    df = df.copy()
    df["Dia_Semana"] = df["Data"].dt.dayofweek.map(DIAS_PT)
    fer = feriados(range(df["Data"].dt.year.min(), df["Data"].dt.year.max() + 1))
    df["Feriado"] = df["Data"].dt.date.map(fer).fillna("")

    df["Alerta_Nao_Util"] = df["Data"].dt.dayofweek.isin([5, 6]) | (df["Feriado"] != "")
    df["Alerta_Alcada"] = (df["Valor"] >= p.alcada * (1 - p.margem_pct / 100)) & (df["Valor"] < p.alcada)
    df["Alerta_Reincidencia"] = df["Alerta_Alcada"] & (df.groupby("Usuario")["Alerta_Alcada"].transform("sum") >= 2)
    df = df.join(_duplicidades(df, p)).join(_fracionamento(df, p))

    if "Aprovador" in df.columns:
        df["Alerta_SoD"] = (df["Aprovador"] != "") & (df["Aprovador"].str.lower() == df["Usuario"].str.lower())
    else:
        df["Alerta_SoD"] = False

    df["Alerta_Fora_Horario"] = False
    if "Hora_Num" in df.columns:
        fora = df["Hora_Num"].notna() & ((df["Hora_Num"] < p.hora_ini) | (df["Hora_Num"] >= p.hora_fim))
        if "Tipo_Lancamento" in df.columns:  # lançamentos automáticos rodam à noite por natureza
            fora &= df["Tipo_Lancamento"].str.lower().str.startswith("man")
        df["Alerta_Fora_Horario"] = fora

    df["Alerta_Valor_Redondo"] = (df["Valor"] >= p.redondo_min) & (df["Valor"].round(2) % 1000 == 0)

    lv = np.log1p(df["Valor"].clip(lower=0))
    media = lv.groupby(df["Conta_Contabil"]).transform("mean")
    desvio = lv.groupby(df["Conta_Contabil"]).transform("std").replace(0, np.nan)
    df["Z_Valor_Conta"] = ((lv - media) / desvio).fillna(0).round(2)
    df["Alerta_Outlier_Valor"] = df["Z_Valor_Conta"] > p.z_corte
    return df


# ---------------------------------------------------------------------------
# IA (Isolation Forest) e Score
# ---------------------------------------------------------------------------
def _features(df: pd.DataFrame) -> pd.DataFrame:
    X = pd.DataFrame(index=df.index)
    X["log_valor"] = np.log1p(df["Valor"].clip(lower=0))
    X["dia_semana"] = df["Data"].dt.dayofweek
    X["nao_util"] = df["Alerta_Nao_Util"].astype(int)
    X["z_valor_conta"] = df["Z_Valor_Conta"]
    n_uc = df.groupby(["Usuario", "Conta_Contabil"])["Valor"].transform("size")
    n_u = df.groupby("Usuario")["Valor"].transform("size")
    X["freq_usuario_conta"] = n_uc / n_u  # raridade da combinação usuário x conta
    if "Hora_Num" in df.columns:
        X["hora"] = df["Hora_Num"].fillna(12)
    if "Tipo_Lancamento" in df.columns:
        X["manual"] = df["Tipo_Lancamento"].str.lower().str.startswith("man").astype(int)
    if "Fornecedor" in df.columns:
        X["raridade_fornecedor"] = -np.log(df.groupby("Fornecedor")["Valor"].transform("size") / len(df))
    return X.join(pd.get_dummies(df["Conta_Contabil"], prefix="conta").astype(int))


def _percentil_ia(df: pd.DataFrame, p: Parametros) -> pd.Series:
    X = _features(df)
    modelo = IsolationForest(n_estimators=200, contamination=p.contaminacao, random_state=42).fit(X)
    return pd.Series(-modelo.score_samples(X), index=df.index).rank(pct=True) * 100


def _motivo(r: dict, p: Parametros) -> str:
    m = []
    if r["Alerta_SoD"]:
        m.append("Aprovador igual ao lançador (violação de segregação de funções)")
    if r["Alerta_Duplicidade"]:
        m.append("Possível duplicidade: " + r["Detalhe_Duplicidade"])
    if r["Alerta_Fracionamento"]:
        m.append("Possível fracionamento: " + r["Detalhe_Fracionamento"])
    if r["Alerta_Alcada"]:
        m.append(f"Valor {brl(r['Valor'])} logo abaixo da alçada de {brl(p.alcada)}")
    if r["Alerta_Reincidencia"]:
        m.append("Usuário repete lançamentos logo abaixo da alçada (padrão recorrente)")
    if r["Alerta_Nao_Util"]:
        m.append(f"Lançado em feriado ({r['Feriado']})" if r["Feriado"] else f"Lançado em {r['Dia_Semana']}")
    if r["Alerta_Fora_Horario"]:
        m.append(f"Lançamento manual fora do expediente ({r.get('Hora', '')})")
    if r["Alerta_Valor_Redondo"]:
        m.append("Valor redondo")
    if r["Alerta_Outlier_Valor"]:
        m.append(f"Valor atípico para a conta (z={r['Z_Valor_Conta']:.1f})")
    if r["Pontos_IA"] > 0 and r["Percentil_IA"] >= 90:
        m.append(f"Padrão atípico segundo a IA (percentil {r['Percentil_IA']:.0f})")
    return "; ".join(m)


def executar(df: pd.DataFrame, p: Parametros) -> pd.DataFrame:
    df = aplicar_regras(df, p)
    df["Percentil_IA"] = _percentil_ia(df, p).round(1)
    df["Pontos_IA"] = (p.peso_ia * ((df["Percentil_IA"] - 80) / 20).clip(0, 1)).round(1)
    pontos_regras = sum(df[c].astype(int) * w for c, w in p.pesos.items())
    df["Score_Risco"] = (pontos_regras + df["Pontos_IA"]).clip(upper=100).round(1)

    df["Nivel_Risco"] = "Baixo"
    for nome, minimo in reversed(NIVEIS):
        df.loc[df["Score_Risco"] >= minimo, "Nivel_Risco"] = nome
    df["Acao_Sugerida"] = df["Nivel_Risco"].map(ACOES)
    df["Qtd_Alertas"] = df[list(ROTULOS)].sum(axis=1)

    df["Motivos"] = ""
    mask = (df["Qtd_Alertas"] > 0) | (df["Pontos_IA"] > 0)
    if mask.any():
        df.loc[mask, "Motivos"] = [_motivo(r, p) for r in df.loc[mask].to_dict("records")]
    return df


# ---------------------------------------------------------------------------
# Lei de Benford e perfil de usuários
# ---------------------------------------------------------------------------
def benford(valores: pd.Series) -> tuple[pd.DataFrame, float, str, int]:
    v = valores[valores >= 1]
    n = len(v)
    esp = pd.Series({d: np.log10(1 + 1 / d) for d in range(1, 10)})
    if n == 0:
        return pd.DataFrame(), np.nan, "Sem dados", 0
    primeiro = (v / 10 ** np.floor(np.log10(v))).astype(int).clip(1, 9)
    obs = primeiro.value_counts(normalize=True).reindex(range(1, 10)).fillna(0)
    mad = float((obs - esp).abs().mean())
    classe = ("Conformidade alta" if mad < 0.006 else "Conformidade aceitável" if mad < 0.012
              else "Conformidade marginal" if mad < 0.015 else "Não conforme")
    tab = pd.DataFrame({"Digito": range(1, 10), "Observado (%)": (obs.values * 100).round(2),
                        "Esperado (%)": (esp.values * 100).round(2)})
    tab["Desvio (p.p.)"] = (tab["Observado (%)"] - tab["Esperado (%)"]).round(2)
    return tab, mad, classe, n


def benford_por_conta(df: pd.DataFrame) -> pd.DataFrame:
    linhas = []
    for conta, g in df.groupby("Conta_Contabil"):
        _, mad, classe, n = benford(g["Valor"])
        linhas.append({"Conta_Contabil": conta, "Qtd": n, "MAD": round(mad, 4) if n else np.nan,
                       "Resultado": classe if n >= 300 else "Amostra < 300 (inconclusivo)"})
    return pd.DataFrame(linhas)


def perfil_usuarios(df: pd.DataFrame) -> pd.DataFrame:
    g = df.groupby("Usuario")
    out = pd.DataFrame({
        "Lançamentos": g.size(),
        "Valor total": g["Valor"].sum().round(2),
        "Críticos+Altos": g["Nivel_Risco"].apply(lambda s: s.isin(["Crítico", "Alto"]).sum()),
        "Score médio": g["Score_Risco"].mean().round(1),
        "Score máximo": g["Score_Risco"].max(),
    })
    out["% Críticos+Altos"] = (out["Críticos+Altos"] / out["Lançamentos"] * 100).round(1)
    return out.sort_values(["Críticos+Altos", "Score médio"], ascending=False).reset_index()


# ---------------------------------------------------------------------------
# Evidência e exportação (papel de trabalho)
# ---------------------------------------------------------------------------
def evidencia(raw: bytes, nome_arquivo: str, p: Parametros, n_linhas: int) -> dict:
    h = hashlib.sha256(raw).hexdigest()
    params_json = json.dumps(asdict(p), sort_keys=True, ensure_ascii=False)
    exec_id = hashlib.sha256((h + params_json + VERSAO).encode()).hexdigest()[:12].upper()
    agora = datetime.now(timezone(timedelta(hours=-3))).strftime("%d/%m/%Y %H:%M:%S (BRT)")
    return {"ID da execução": exec_id, "Arquivo": nome_arquivo, "SHA-256 do arquivo": h,
            "Linhas analisadas": n_linhas, "Executado em": agora, "Versão do motor": VERSAO,
            "Parâmetros": params_json}


def colunas_saida(df: pd.DataFrame) -> list[str]:
    base = ["ID_Transacao", "Data", "Dia_Semana", "Hora", "Usuario", "Aprovador", "Conta_Contabil", "Fornecedor",
            "Numero_NF", "Tipo_Lancamento", "Valor", "Score_Risco", "Nivel_Risco", "Acao_Sugerida", "Motivos",
            *ROTULOS, "Percentil_IA", "Z_Valor_Conta"]
    return [c for c in base if c in df.columns]


STATUS = ["Novo", "Em análise", "Procedente", "Improcedente", "Falso positivo"]


def exportar_excel(df: pd.DataFrame, evid: dict, benf: pd.DataFrame, triagem: dict | None = None) -> bytes:
    triagem = triagem or {}
    saida = df[colunas_saida(df)].copy()
    saida["Data"] = saida["Data"].dt.date
    alertas = saida[saida["Nivel_Risco"].isin(["Crítico", "Alto", "Médio"])].sort_values("Score_Risco", ascending=False)

    tri = alertas[["ID_Transacao", "Usuario", "Valor", "Score_Risco", "Nivel_Risco", "Motivos"]].copy()
    tri["Status"] = tri["ID_Transacao"].map(lambda i: triagem.get(i, {}).get("Status", "Novo"))
    tri["Responsavel"] = tri["ID_Transacao"].map(lambda i: triagem.get(i, {}).get("Responsavel", ""))
    tri["Parecer"] = tri["ID_Transacao"].map(lambda i: triagem.get(i, {}).get("Parecer", ""))

    resumo = pd.DataFrame({"Nível": ["Crítico", "Alto", "Médio", "Baixo"]})
    resumo["Transações"] = resumo["Nível"].map(df["Nivel_Risco"].value_counts()).fillna(0).astype(int)
    resumo["Valor (R$)"] = resumo["Nível"].map(df.groupby("Nivel_Risco")["Valor"].sum()).fillna(0).round(2)

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        pd.DataFrame({"Item": list(evid), "Valor": [str(v) for v in evid.values()]}).to_excel(xw, sheet_name="Evidencia", index=False)
        resumo.to_excel(xw, sheet_name="Resumo", index=False)
        alertas.to_excel(xw, sheet_name="Alertas", index=False)
        tri.to_excel(xw, sheet_name="Triagem", index=False)
        benf.to_excel(xw, sheet_name="Benford", index=False)
        saida.to_excel(xw, sheet_name="Base_Completa", index=False)
        from openpyxl.styles import Font, PatternFill
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F3A5F")
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                largura = max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:200])
                ws.column_dimensions[col[0].column_letter].width = min(max(10, largura + 2), 70)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# API de alto nível (usada pelo app, pela CLI e por agentes)
# ---------------------------------------------------------------------------
def analisar_arquivo(raw: bytes, nome: str = "arquivo.csv", p: Parametros | None = None) -> dict:
    p = p or Parametros()
    df, avisos = validar(ler_csv_bytes(raw))
    res = executar(df, p)
    tab, mad, classe, n = benford(res["Valor"])
    return {"df": res, "avisos": avisos, "benford": tab, "benford_mad": mad, "benford_classe": classe,
            "benford_n": n, "benford_contas": benford_por_conta(res), "evidencia": evidencia(raw, nome, p, len(res))}


def main() -> None:
    ap = argparse.ArgumentParser(description="Motor de Auditoria Contínua")
    ap.add_argument("arquivo", help="CSV de lançamentos")
    ap.add_argument("-o", "--saida", default="resultado_auditoria.xlsx")
    ap.add_argument("--alcada", type=float, default=10000.0)
    a = ap.parse_args()
    raw = open(a.arquivo, "rb").read()
    r = analisar_arquivo(raw, a.arquivo, Parametros(alcada=a.alcada))
    open(a.saida, "wb").write(exportar_excel(r["df"], r["evidencia"], r["benford"]))
    print(f"Execução {r['evidencia']['ID da execução']} | {len(r['df'])} lançamentos")
    print(r["df"]["Nivel_Risco"].value_counts().to_string())
    print(f"Papel de trabalho salvo em {a.saida}")


if __name__ == "__main__":
    main()